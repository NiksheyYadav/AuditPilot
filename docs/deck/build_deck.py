"""Builds AuditPilot_IT2_Presentation.pptx (python-pptx). Screenshots come from docs/screenshots (real captures)."""
import os, json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, '..', 'screenshots')
TMP = os.path.join(HERE, '_crops'); os.makedirs(TMP, exist_ok=True)
OUT = os.path.join(HERE, '..', '..', 'AuditPilot_IT2_Presentation.pptx')
TOTAL = 12
T = json.load(open(os.path.join(HERE, 'facts.json')))   # real numbers measured from the running app and npm test

def C(h): return RGBColor.from_string(h)
INK, MUTED, BG, WHITE, LINE = '1C1736', '6B6788', 'FAF7FF', 'FFFFFF', 'E6E2F3'
AMETHYST, FUCHSIA, EMERALD, AMBER, RUBY, SAPPHIRE = '6D28D9', 'C026D3', '059669', 'D97706', 'E11D48', '2563EB'
SOFT = {AMETHYST: 'EDE9FE', FUCHSIA: 'FAE8FF', EMERALD: 'D1FAE5', AMBER: 'FEF3C7', RUBY: 'FFE4E6', SAPPHIRE: 'DBEAFE'}
F, FB = 'Segoe UI', 'Segoe UI Semibold'

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
count = [0]

def bg(s, dark=False):
    f = s.background.fill
    if dark:
        f.gradient(); f.gradient_angle = 35
        f.gradient_stops[0].color.rgb = C('2E1065'); f.gradient_stops[1].color.rgb = C('A21CAF')
    else:
        f.solid(); f.fore_color.rgb = C(BG)

def shape(s, kind, x, y, w, h, fill=WHITE, line=None, radius=0.1, lw=1.0):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    if line: sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
    else: sh.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    return sh

def card(s, x, y, w, h, fill=WHITE, line=LINE, radius=0.06): return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius)

def grad(sh, a='5B21B6', b='C026D3'):
    sh.fill.gradient(); sh.fill.gradient_angle = 0
    sh.fill.gradient_stops[0].color.rgb = C(a); sh.fill.gradient_stops[1].color.rgb = C(b)

def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.05, spacing=1.0):
    if isinstance(paras, tuple): paras = [paras]
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, (t, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        if i: p.space_before = Pt(size * 0.45)
        r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = False
        r.font.name = FB if bold else F; r.font.color.rgb = C(color)
    return tb

def boxed(s, x, y, w, h, title, sub=None, color=AMETHYST, fill=None, tsize=15, ssize=11.5, dark_text=False):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill or color, None, 0.12)
    tc = INK if dark_text else WHITE
    text(s, x, y, w, h, [(title, tsize, True, tc)] + ([(sub, ssize, False, tc)] if sub else []), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.08)

def arrow(s, x1, y1, x2, y2, color=MUTED, w=2.25):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = C(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln(); t = etree.SubElement(ln, qn('a:tailEnd')); t.set('type', 'triangle'); t.set('w', 'med'); t.set('len', 'med')

def num(s, x, y, n, color, d=0.5, size=14):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, d, d, color, None, 0.25)
    text(s, x, y, d, d, (n, size, True, WHITE), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)

def new(label, title, note, notes, dark=False):
    """Slide chrome follows the reference deck: TRAINING / NN kicker, sub-label, title, N / total counter, status line."""
    s = prs.slides.add_slide(BLANK); count[0] += 1; bg(s, dark)
    if not dark:
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 0.62, 0.09, 0.95, FUCHSIA, None, 0.5)
        text(s, 0.92, 0.36, 4, 0.3, (f'TRAINING  /  {count[0]:02d}', 11, True, FUCHSIA))
        text(s, 0.92, 0.6, 9, 0.3, (label.upper(), 10.5, True, MUTED))
        text(s, 0.92, 0.88, 11.6, 0.75, (title, 30, True, INK))
        text(s, 0.7, 7.02, 10.8, 0.3, (note, 10.5, False, MUTED))
        text(s, 11.6, 7.02, 1.05, 0.3, (f'{count[0]}  /  {TOTAL}', 10.5, True, MUTED), PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s

def crop(src, box, name):
    im = Image.open(os.path.join(SHOTS, src)).convert('RGB')
    im.crop(box).save(os.path.join(TMP, name), quality=92); return os.path.join(TMP, name)

NT, NK = T['tests_total'], T['nik']

# ================================================================ 1 Title
s = new('', '', '', (
    "Namaste, main Nikshey Yadav hoon, SGT University se B.Tech CSE Data Science, 7th semester.\n"
    "Meri training AccioJobs Pvt. Ltd. mein thi, aur ye mera project hai: AuditPilot.\n"
    "AuditPilot ek AI website audit agent hai: tools facts nikalte hain, local open-source LLM unhe explain karta hai.\n"
    "Slide par jo 54 aur 96 hai, wo mere project ke do demo pages ka asli audit score hai."), dark=True)
text(s, 0.9, 0.8, 7, 0.4, ('ACCIOJOBS PVT. LTD.  ·  INDUSTRIAL TRAINING PROJECT (IT-2)', 13, True, 'F5D0FE'))
text(s, 0.9, 1.5, 7.5, 1.4, ('AuditPilot', 72, True, WHITE))
text(s, 0.9, 2.95, 7.2, 1.3, ('An AI website audit agent. Tools compute the facts; a local open-source LLM explains them.', 22, False, 'EDE9FE'), spacing=1.05)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 4.45, 0.9, 0.06, 'F0ABFC', None, 0.5)
text(s, 0.9, 4.6, 7.5, 1.9, [('Nikshey Yadav  ·  Roll No. 231302221', 20, True, WHITE),
    ('B.Tech CSE (Data Science)  ·  7th Semester  ·  SGT University', 15, False, 'EDE9FE'),
    ('Training: AccioJobs Pvt. Ltd.  ·  8 August to 18 March', 15, False, 'EDE9FE')])
text(s, 0.9, 6.7, 7.5, 0.4, ('REACT  ·  EXPRESS  ·  MONGODB  ·  OLLAMA', 12, True, 'F0ABFC'))
card(s, 8.9, 1.45, 3.7, 4.3, WHITE, None, 0.07)
text(s, 8.9, 1.65, 3.7, 0.4, ('SAME PAGE, BEFORE AND AFTER FIXES', 10.5, True, MUTED), PP_ALIGN.CENTER)
text(s, 9.0, 2.1, 1.6, 1.3, (str(T['bad']), 60, True, AMBER), PP_ALIGN.CENTER)
text(s, 10.9, 2.1, 1.6, 1.3, (str(T['good']), 60, True, EMERALD), PP_ALIGN.CENTER)
arrow(s, 10.45, 2.75, 10.95, 2.75, MUTED, 3)
text(s, 9.0, 3.25, 1.6, 0.4, ('/demo/bad', 12, True, MUTED), PP_ALIGN.CENTER)
text(s, 10.9, 3.25, 1.6, 0.4, ('/demo/good', 12, True, MUTED), PP_ALIGN.CENTER)
text(s, 9.1, 4.0, 3.5, 1.5, [(f"+{T['good'] - T['bad']} points", 24, True, AMETHYST), ('Real scores from the running app, computed with no LLM', 12, False, MUTED)], PP_ALIGN.CENTER)

# ================================================================ 2 Context + roadmap
s = new('Setting the context', 'Training Context & Learning Roadmap',
    'Training covered the first five; Express, MongoDB and JWT were learned with AI-assisted development (Claude Code).', (
    "Training mein humne HTML, CSS, JavaScript, React, AI aur AI agents seekhe, Node.js thoda, aur MCP ke basics.\n"
    "Ye pehli panch cards training ka roadmap hain.\n"
    "Express, MongoDB aur JWT training mein cover nahi hue the. Inhe maine Claude Code ke saath AI-assisted development se seekha aur is project mein use kiya.\n"
    "Main ye openly bata raha hoon, aur har feature explain kar sakta hoon."))
text(s, 0.7, 1.85, 6, 0.3, ('COVERED IN TRAINING', 11, True, EMERALD))
road = [('HTML + CSS', 'structure and layout'), ('JavaScript', 'logic, events, data'), ('React', 'components, state'), ('AI + AI agents', 'LLMs, tool calling'), ('Node.js + MCP', 'a little, basics')]
for i, (a, b) in enumerate(road):
    x = 0.7 + i * 2.4
    card(s, x, 2.2, 2.2, 1.75); num(s, x + 0.15, 2.35, f'0{i + 1}', EMERALD)
    text(s, x + 0.12, 3.0, 2.0, 0.9, [(a, 15, True, INK), (b, 11.5, False, MUTED)])
    if i < 4: shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 2.2, 3.05, 0.2, 0.04, LINE, None, 0.5)
text(s, 0.7, 4.3, 8, 0.3, ('EXTENDED BEYOND TRAINING (AI-ASSISTED DEVELOPMENT WITH CLAUDE CODE)', 11, True, AMETHYST))
for i, (a, b) in enumerate([('Express', 'REST API, SSE streaming, middleware'), ('MongoDB', 'Mongoose models and indexes'), ('JWT', 'login and protected routes')]):
    x = 0.7 + i * 4.05
    card(s, x, 4.65, 3.85, 1.25, SOFT[AMETHYST], None); num(s, x + 0.15, 4.85, f'0{i + 6}', AMETHYST)
    text(s, x + 0.8, 4.8, 2.95, 0.95, [(a, 17, True, INK), (b, 12, False, MUTED)], anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 6.2, 11.93, 0.55, SOFT[AMBER], None, 0.3)
text(s, 0.9, 6.2, 11.5, 0.55, ('Current foundation:  HTML  ·  CSS  ·  JavaScript  ·  React  +  Express, MongoDB, JWT  +  a local LLM', 14, True, '92400E'), anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 3 Problem -> solution
s = new('Problem to solution', 'Project Overview: AuditPilot',
    'Project status: implemented, tested and demo-ready. Scores are deterministic; the AI only explains them.', (
    "Problem ye hai ki website quality ke liye alag alag tools chalane padte hain: SEO, accessibility, security, performance.\n"
    "Aur report mein sirf raw data hota hai, ye samajh nahi aata ki pehle kya fix karein.\n"
    "AuditPilot URL leta hai, 29 checks chalata hai, score aur evidence deta hai. Phir local AI batata hai ki kya pehle fix karna hai.\n"
    "Objective: website quality ko samajhna aur fix karna aasan banana."))
card(s, 0.7, 1.95, 5.2, 3.75); shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 1.95, 0.12, 3.75, RUBY, None, 0.5)
text(s, 1.1, 2.1, 4.6, 0.5, ('THE PROBLEM', 13, True, RUBY))
text(s, 1.1, 2.6, 4.6, 3.0, [('Quality checks are manual', 19, True, INK), ('A different tool for SEO, accessibility, security and speed.', 14, False, MUTED),
    ('Reports are raw data', 19, True, INK), ('Hundreds of lines, no clear priority.', 14, False, MUTED),
    ('Cloud AI is not always allowed', 19, True, INK), ('Site data may be private.', 14, False, MUTED)])
arrow(s, 6.0, 3.8, 7.0, 3.8, FUCHSIA, 4)
card(s, 7.1, 1.95, 5.53, 3.75); shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.1, 1.95, 0.12, 3.75, EMERALD, None, 0.5)
text(s, 7.5, 2.1, 5.0, 0.5, ('THE SOLUTION: AUDITPILOT', 13, True, EMERALD))
text(s, 7.5, 2.6, 5.0, 3.0, [('One URL, 29 checks', 19, True, INK), ('Scores and evidence in milliseconds.', 14, False, MUTED),
    ('A local AI explains', 19, True, INK), ('Summary, fix plan and chat on your laptop.', 14, False, MUTED),
    ('Compare, schedule, render JS', 19, True, INK), ('Before and after, timed audits, client-side apps.', 14, False, MUTED)])
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 5.95, 11.93, 0.8, SOFT[AMETHYST], None, 0.25)
text(s, 0.95, 5.95, 11.4, 0.8, ('Objective: make website quality easy to measure, understand and fix, without sending data to the cloud.', 16, True, AMETHYST), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# ================================================================ 4 Tech stack
s = new('What each technology does', 'Technology Stack & Tools',
    f'Everything runs locally: no paid or cloud AI APIs. {NT} automated tests cover the server.', (
    "Frontend React 19 aur Vite hai, charts ke liye Recharts.\n"
    "Backend Express 5 hai, MongoDB Mongoose ke saath, login ke liye JWT aur bcrypt.\n"
    "AI ke liye Ollama par open-source model qwen2.5:3b chalta hai: private, free, aur offline kaam karta hai.\n"
    f"MCP SDK se same tools Claude Code mein bhi chalte hain, aur Vitest se {NT} tests likhe hain."))
tiles = [('React 19 + Vite', 'fast component UI', AMETHYST, 'INTERFACE'), ('Express 5', 'REST API and SSE streaming', FUCHSIA, 'SERVER'), ('MongoDB + Mongoose', 'nested audit documents', EMERALD, 'DATA'),
         ('JWT + bcrypt', 'stateless login, hashed passwords', SAPPHIRE, 'AUTH'), ('cheerio + playwright-core', 'HTML parsing, optional JS render', AMBER, 'AUDIT'), ('Recharts + zod', 'trend charts, input validation', RUBY, 'UI + SAFETY'),
         ('MCP SDK', 'same tools in Claude Code', SAPPHIRE, 'INTEGRATION'), ('Vitest + supertest', f'{NT} automated tests', EMERALD, 'QUALITY')]
for i, (a, b, col, tag) in enumerate(tiles):
    x, y = 0.7 + (i % 4) * 3.0, 1.95 + (i // 4) * 1.55
    card(s, x, y, 2.85, 1.4); shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 2.85, 0.12, col, None, 0.5)
    text(s, x + 0.1, y + 0.22, 2.65, 1.15, [(a, 15, True, INK), (tag, 9.5, True, col), (b, 11.5, False, MUTED)])
g = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 5.1, 11.93, 1.7, AMETHYST, None, 0.08); grad(g)
text(s, 1.0, 5.2, 5.6, 1.5, [('Ollama + qwen2.5:3b', 24, True, WHITE), ('Open-source LLM running locally', 14, False, 'F5D0FE')], anchor=MSO_ANCHOR.MIDDLE)
for i, w in enumerate(['Private', 'Free', 'Offline']):
    boxed(s, 7.0 + i * 1.9, 5.62, 1.75, 0.65, w, color=WHITE, fill='FFFFFF', tsize=16, dark_text=True)

# ================================================================ 5 Architecture
s = new('How the app works', 'System Architecture & Workflow',
    'Also: a scheduler runs timed audits, and an optional headless browser renders JavaScript, both behind the same SSRF guard.', (
    "Ye architecture hai. Browser mein React UI Express API se baat karta hai.\n"
    "Hari line deterministic path hai: Express audit engine ko call karta hai, jo SSRF guard ke saath target website fetch karta hai. Isme LLM nahi hai.\n"
    "Orange line AI path hai: agent loop queue ke through local Ollama model se baat karta hai.\n"
    "Same tool registry MCP server ke through Claude Code ko bhi milti hai, aur sab data MongoDB mein save hota hai."))
C1, C2, C3, C4 = 0.7, 3.45, 6.45, 9.9
boxed(s, C1, 3.55, 2.2, 1.0, 'React UI', 'Vite, Recharts', AMETHYST)
boxed(s, C2, 3.55, 2.55, 1.0, 'Express API', 'auth, audits, chat (SSE)', FUCHSIA)
boxed(s, C2, 1.85, 2.55, 0.95, 'MongoDB', 'users, audits, schedules', SAPPHIRE)
boxed(s, C3, 1.85, 2.95, 0.95, 'Audit engine', '29 checks + scoring', EMERALD)
boxed(s, C4, 1.85, 2.7, 0.95, 'Target website', 'SSRF-guarded fetch', '6B6788')
boxed(s, C3, 3.55, 2.95, 1.0, 'Agent loop', 'max 5 steps, zod-validated', AMBER)
boxed(s, C4, 3.55, 2.7, 1.0, 'Ollama LLM', 'qwen2.5:3b via queue', RUBY)
boxed(s, C3, 5.25, 2.95, 0.95, 'Tool registry', '5 tools, one source', AMBER)
boxed(s, C2, 5.25, 2.55, 0.95, 'MCP server', 'stdio', SAPPHIRE)
boxed(s, C1, 5.25, 2.2, 0.95, 'Claude Code', 'or Claude Desktop', '6B6788')
arrow(s, 2.9, 4.05, 3.45, 4.05, INK); arrow(s, 5.1, 3.55, 5.1, 2.8, SAPPHIRE); arrow(s, 6.0, 3.7, 6.45, 2.5, EMERALD, 3)
arrow(s, 9.4, 2.32, 9.9, 2.32, EMERALD, 3); arrow(s, 6.0, 4.05, 6.45, 4.05, AMBER, 3); arrow(s, 9.4, 4.05, 9.9, 4.05, AMBER, 3)
arrow(s, 7.9, 4.55, 7.9, 5.25, AMBER, 3); arrow(s, 2.9, 5.72, 3.45, 5.72, MUTED); arrow(s, 6.0, 5.72, 6.45, 5.72, MUTED)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 6.55, 0.3, 0.14, EMERALD, None, 0.5); text(s, 1.05, 6.45, 4.2, 0.35, ('Deterministic path: no LLM involved', 12, False, INK))
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 5.4, 6.55, 0.3, 0.14, AMBER, None, 0.5); text(s, 5.75, 6.45, 4.4, 0.35, ('AI path: local, optional, queued', 12, False, INK))

# ================================================================ 6 Agent loop
s = new('How the AI agent works', 'Tools compute the facts. The LLM only explains.',
    'Streaming together with tool calls was tested on the installed Ollama; the answer arrives token by token.', (
    "Core principle ye hai: tools facts compute karte hain, LLM sirf explain karta hai.\n"
    "Chat mein current audit ka data pehle se prompt mein hota hai, to zyada sawal bina tool ke solve ho jaate hain.\n"
    "Agar zarurat ho to model tool maangta hai, server zod se arguments validate karke tool chalata hai, aur result wapas model ko deta hai.\n"
    "Guardrails: maximum 5 steps, galat arguments par ek retry, user sirf apne audits dekh sakta hai, aur ek time par ek hi LLM call."))
boxed(s, 0.7, 2.0, 2.3, 1.0, 'Your question', 'plus audit data preloaded', AMETHYST)
boxed(s, 3.6, 2.0, 2.3, 1.0, 'Local LLM', 'qwen2.5:3b, temp 0.1', RUBY)
shape(s, MSO_SHAPE.DIAMOND, 6.5, 1.85, 2.1, 1.3, AMBER, None)
text(s, 6.5, 1.85, 2.1, 1.3, ('Tool needed?', 14, True, WHITE), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
boxed(s, 9.4, 2.0, 3.2, 1.0, 'Streamed answer', 'tokens arrive live (SSE)', EMERALD)
boxed(s, 6.2, 4.0, 2.7, 1.0, 'Run tool', 'zod-validated arguments', AMBER)
boxed(s, 3.0, 4.0, 2.7, 1.0, 'Result to LLM', 'name, args, ms saved as trace', FUCHSIA)
arrow(s, 3.0, 2.5, 3.6, 2.5, INK); arrow(s, 5.9, 2.5, 6.5, 2.5, INK); arrow(s, 8.6, 2.5, 9.4, 2.5, EMERALD, 3)
text(s, 8.55, 2.0, 0.9, 0.35, ('no', 12, True, EMERALD), PP_ALIGN.CENTER)
arrow(s, 7.55, 3.15, 7.55, 4.0, AMBER, 3); text(s, 7.65, 3.4, 0.8, 0.35, ('yes', 12, True, AMBER))
arrow(s, 6.2, 4.5, 5.7, 4.5, AMBER, 3); arrow(s, 4.35, 4.0, 4.75, 3.0, FUCHSIA, 3)
text(s, 0.7, 3.45, 2.5, 0.35, ('Max 5 steps per question', 12, False, MUTED))
text(s, 0.7, 5.45, 6, 0.4, ('GUARDRAILS', 12, True, FUCHSIA))
rails = [('Max 5 steps', AMETHYST), ('Bad args: retry once', AMBER), ('Own audits only', EMERALD), ('SSRF-guarded run_audit', RUBY), ('1 LLM call at a time', SAPPHIRE), ('Short refs, no id leaks', FUCHSIA)]
for i, (r, col) in enumerate(rails):
    boxed(s, 0.7 + (i % 3) * 4.0, 5.85 + (i // 3) * 0.5, 3.8, 0.4, r, color=col, fill=SOFT[col], tsize=12.5, dark_text=True)

# ================================================================ 7 ER diagram
s = new('Data model', 'Three collections, one nested document',
    'Findings, chat messages and tool calls are embedded in the audit, so one query loads a whole report.', (
    "Database MongoDB mein hai. Teen collections hain: User, Audit aur Schedule.\n"
    "Ek user ke bahut saare audits aur schedules hote hain. Audit mein index user, url aur createdAt par hai.\n"
    "Findings, chat messages aur tool calls Audit document ke andar embedded hain, isliye ek query mein poori report mil jaati hai.\n"
    "Schedule collection timed audits ke liye hai: kitne ghante baad dobara audit karna hai."))
def entity(x, y, w, h, name, fields, col):
    card(s, x, y, w, h); shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.5, col, None, 0.25)
    text(s, x + 0.1, y, w - 0.2, 0.5, (name, 15, True, WHITE), anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.12, y + 0.58, w - 0.2, h - 0.6, [(f, 11.5, False, INK) for f in fields], spacing=0.95)
entity(0.7, 1.95, 2.9, 1.75, 'User', ['_id, name', 'email (unique)', 'passwordHash (bcrypt)'], AMETHYST)
entity(0.7, 4.45, 2.9, 2.05, 'Schedule', ['user -> User, url', 'everyHours, render', 'nextRunAt, lastRunAt', 'lastAuditId, lastError'], SAPPHIRE)
entity(4.6, 1.95, 3.6, 4.55, 'Audit', ['_id', 'user  ->  User', 'url, finalUrl, status, rendered', 'scores {overall, seo, accessibility,', '            security, performance}', 'summary, fixPlan', 'aiStatus: pending|ready|skipped|failed', 'llm {provider, model, mock}', 'durationMs, createdAt', 'index (user, url, createdAt)'], FUCHSIA)
entity(9.3, 1.95, 3.33, 1.85, 'Finding (embedded, 29)', ['id, category, severity', 'title, evidence', 'recommendation'], EMERALD)
entity(9.3, 4.15, 3.33, 1.5, 'ChatMessage (embedded)', ['role, content, note', 'toolCalls[]'], AMBER)
card(s, 9.3, 5.9, 3.33, 0.7, SOFT[RUBY], None); text(s, 9.4, 5.9, 3.1, 0.7, [('ToolCall: name, args, resultSummary, ms', 11.5, True, '9F1239')], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, 3.6, 2.8, 4.6, 2.8, INK, 2.5); text(s, 3.55, 2.35, 1.1, 0.35, ('1 : many', 11, True, MUTED), PP_ALIGN.CENTER)
arrow(s, 2.15, 3.7, 2.15, 4.45, SAPPHIRE, 2.5); text(s, 2.25, 3.9, 1.1, 0.35, ('1 : many', 11, True, MUTED))
arrow(s, 8.2, 2.9, 9.3, 2.9, EMERALD, 2.5); text(s, 8.15, 2.45, 1.2, 0.35, ('embeds', 11, True, MUTED), PP_ALIGN.CENTER)
arrow(s, 8.2, 4.9, 9.3, 4.9, AMBER, 2.5); text(s, 8.15, 4.45, 1.2, 0.35, ('embeds', 11, True, MUTED), PP_ALIGN.CENTER)
arrow(s, 10.95, 5.65, 10.95, 5.9, RUBY, 2.5)

# ================================================================ 8 Features
s = new('What users can do', 'Functional Modules & Key Features',
    f"Screenshots are real captures of the running app with the local model (qwen2.5:3b): https://nikshey.xyz and the /demo/site crawl.", (
    f"Ye saare screenshots running app se Playwright se liye hain, asli local model qwen2.5:3b ke saath, aur real website nikshey.xyz ke audit par. Koi mock ya seeded text nahi hai.\n"
    f"Pehla: animated score gauge, chaar category scores aur AI summary. Doosra: har finding mein evidence aur Nginx ya Express fix snippet.\n"
    "Teesra: agent chat, aur collapsible tool trace jisme tool ka naam, args aur time dikhta hai.\n"
    f"Chautha screenshot poori website ka crawl hai: demo site ke {T['site_audit']['pages']} pages, site score {T['site_audit']['overall']}, aur cross-page findings jaise duplicate titles aur broken link.\n"
    "Neeche saat modules hain: audit engine, AI summary aur chat, compare aur trend, history aur sign-in, scheduled audits, JavaScript rendering, aur site crawl."))
SA = T['site_audit']
imgs = [(crop('04-report-top.png', (248, 130, 2628, 1452), 'a.png'), 'Audit report', f"nikshey.xyz scored {NK['overall']}: gauge, category scores and the AI summary"),
        (crop('05-findings-fix.png', (248, 140, 1648, 918), 'b.png'), 'Evidence and fix', 'Every finding shows its evidence and a ready fix snippet'),
        (crop('06-chat-tool-trace.png', (0, 395, 940, 917), 'c.png'), 'Agent chat', 'A collapsible tool trace shows the tool, its arguments and the time'),
        (crop('14-site-audit.png', (240, 110, 2640, 1443), 'e.png'), 'Whole-site audit', f"/demo/site: {SA['pages']} pages crawled, site score {SA['overall']}, plus cross-page findings")]
for k, (p, cap, sub) in enumerate(imgs):
    x, y = 0.7 + (k % 2) * 6.05, 1.75 + (k // 2) * 2.35
    card(s, x, y, 5.9, 2.25, WHITE, LINE, 0.03)
    s.shapes.add_picture(p, Inches(x + 0.08), Inches(y + 0.08), width=Inches(3.76), height=Inches(2.09))
    text(s, x + 3.95, y + 0.1, 1.9, 2.05, [(cap, 14, True, INK), (sub, 11, False, MUTED)], anchor=MSO_ANCHOR.MIDDLE)
mods = [('01', 'Audit engine', EMERALD), ('02', 'AI + chat', FUCHSIA), ('03', 'Compare + trend', AMETHYST), ('04', 'History + sign-in', RUBY), ('05', 'Schedules', SAPPHIRE), ('06', 'JS rendering', AMBER), ('07', 'Site crawl', EMERALD)]
for i, (n, a, col) in enumerate(mods):
    x = 0.7 + i * 1.71
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 6.56, 1.65, 0.4, SOFT[col], None, 0.3)
    text(s, x + 0.05, 6.56, 1.6, 0.4, [(f'{n}  {a}', 10.5, True, INK)], anchor=MSO_ANCHOR.MIDDLE)

# ================================================================ 9 Live demo flow
s = new('Live demo sequence', 'Bad page to good page in four minutes',
    f"Live-demo sequence: audit bad -> ask the agent -> audit good -> compare. Real site: nikshey.xyz scored {NK['overall']}/100.", (
    f"Demo ka flow: pehle demo bad page audit karte hain, score {T['bad']} aata hai, turant, bina LLM ke.\n"
    "Findings dekhte hain, aur jab tak local model summary likhta hai, hum findings par baat karte hain. Phir agent se puchte hain: what should I fix first.\n"
    f"Phir demo good page audit karte hain, score {T['good']}, aur compare page par {T['fixed']} issues fixed dikhte hain. Aakhir mein real website nikshey.xyz ka audit, tool trace aur MCP.\n"
    "Ye numbers asli hain: laptop par CPU par model chalta hai, isliye AI answer mein 15 se 40 second lag sakte hain."))
steps = [('1', 'Audit demo (bad)', 'score appears instantly', RUBY), ('2', 'Read the findings', 'while AI writes summary', AMBER), ('3', 'Ask the agent', '"what should I fix first?"', FUCHSIA),
         ('4', 'Audit demo (good)', 'same page, fixed', EMERALD), ('5', 'Compare', f"{T['bad']} to {T['good']}, {T['fixed']} issues fixed", AMETHYST), ('6', 'Real site + MCP', 'nikshey.xyz, tool trace', SAPPHIRE)]
for i, (n, a, b, col) in enumerate(steps):
    x = 0.7 + i * 2.0
    card(s, x, 1.95, 1.88, 2.15); shape(s, MSO_SHAPE.OVAL, x + 0.64, 2.1, 0.6, 0.6, col, None)
    text(s, x + 0.64, 2.1, 0.6, 0.6, (n, 18, True, WHITE), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0)
    text(s, x + 0.05, 2.8, 1.78, 1.3, [(a, 15, True, INK), (b, 12, False, MUTED)], PP_ALIGN.CENTER)
    if i < 5: arrow(s, x + 1.9, 3.0, x + 2.0, 3.0, MUTED, 2)
card(s, 0.7, 4.45, 7.0, 2.35, WHITE)
s.shapes.add_picture(crop('09-compare.png', (250, 300, 2630, 1090), 'd.png'), Inches(1.2), Inches(4.52), width=Inches(6.0), height=Inches(2.0))
text(s, 0.8, 6.5, 6.8, 0.3, (f"Real compare page: {T['bad']} to {T['good']} (+{T['good'] - T['bad']})", 11, True, MUTED), PP_ALIGN.CENTER)
card(s, 8.0, 4.45, 4.63, 2.35, SOFT[AMETHYST], None)
text(s, 8.25, 4.55, 4.2, 2.2, [('Measured on this laptop', 13, True, AMETHYST), ('Audit: tens of milliseconds (local pages)', 13, False, INK), ('AI summary: about 20 to 35 s (CPU)', 13, False, INK), ('Chat answer: about 15 to 40 s', 13, False, INK), (f"Site crawl, {T['site_audit']['pages']} demo pages: about {T['site_audit']['crawl_seconds']:.0f} s", 13, False, INK), ('Scores never wait for the model', 13, True, INK)])

# ================================================================ 10 Security and testing
s = new('Measure functionality, not invented results', 'Security & Testing Checklist',
    f"Counts are from the last npm test run. Live-model behaviour: npm run test:live ({T['live']} tests, opt-in) and npm run check:llm.", (
    "Kyunki hum user ka diya URL fetch karte hain, SSRF guard zaroori hai: sirf http aur https, private aur localhost addresses block, metadata address block, max 3 redirects, 10 second timeout, 2 MB limit.\n"
    "Login JWT se hai, passwords bcrypt se hash hote hain, helmet aur rate limit lage hain, aur ek user doosre ka audit nahi dekh sakta.\n"
    f"Testing mein {NT} tests hain aur sab pass hote hain: {T['audit']} checker aur scoring, {T['ssrf']} SSRF, {T['api']} API aur auth, {T['agent']} agent, {T['features']} nayi features ke (render, schedule, short refs), {T['history']} history aur sign-in ke, {T['site']} site crawl ke (robots.txt, same-origin, cross-page checks, aggregate score), aur {T['mcp']} MCP ke.\n"
    f"Asli local model ke liye alag opt-in suite hai, npm run test:live, jisme {T['live']} tests asli qwen2.5:3b par chalte hain."))
card(s, 0.7, 1.95, 6.2, 4.85)
text(s, 1.0, 2.05, 5.6, 0.45, ('Security layers', 20, True, AMETHYST))
sec = [('SSRF guard', 'http/https only, private + metadata IPs blocked, 3 redirects, 10 s, 2 MB; also applied to every request of the headless browser', RUBY), ('JWT + bcrypt', 'protected routes, hashed passwords', SAPPHIRE),
       ('Ownership checks', "other users' audits return 404 (also for the agent and MCP tools)", EMERALD), ('helmet, CORS, rate limits', 'per-user limits on audits and chat', AMBER), ('zod validation', 'every input and every tool call', FUCHSIA)]
for i, (a, b, col) in enumerate(sec):
    y = 2.6 + i * 0.84
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, y + 0.05, 0.1, 0.62, col, None, 0.5)
    text(s, 1.25, y, 5.5, 0.8, [(a, 14.5, True, INK), (b, 11, False, MUTED)], spacing=0.95)
g = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 1.95, 5.43, 1.85, AMETHYST, None, 0.07); grad(g)
text(s, 7.2, 1.98, 5.43, 1.15, (f'{NT} / {NT}', 54, True, WHITE), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 7.2, 3.1, 5.43, 0.5, ('tests passing  ·  npm test', 15, False, 'F5D0FE'), PP_ALIGN.CENTER)
card(s, 7.2, 4.0, 5.43, 2.8)
rows = [('Checkers, scoring, compare', T['audit']), ('SSRF guard', T['ssrf']), ('Auth, API, ownership', T['api']), ('Agent, summary, SSE chat', T['agent']),
        ('Render, schedules, short refs, links', T['features']), ('History, sessions, sign-up and sign-in', T['history']), ('Site crawl, robots.txt, cross-page checks', T['site']), ('MCP stdio client', T['mcp'])]
for i, (a, b) in enumerate(rows):
    y = 4.05 + i * 0.335
    text(s, 7.4, y, 0.4, 0.33, ('✓', 13, True, EMERALD), anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.8, y, 3.8, 0.33, (a, 12, False, INK), anchor=MSO_ANCHOR.MIDDLE); text(s, 11.6, y, 0.85, 0.33, (str(b), 13, True, AMETHYST), PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

# ================================================================ 11 Learning / challenges / future
s = new('What this project taught me', 'Learning Outcomes & Future Scope',
    'These are the real challenges I hit; future items are proposals, not claims of completion.', (
    "Seekha: poora MERN stack, JWT, SSRF jaise security risks, agent loop aur tool calling, MCP, aur testing.\n"
    "Challenge: chhota 3B local model kabhi tool ke arguments galat bana deta tha, jaise audit ID mangle kar dena. Maine short refs aur title se lookup add kiya.\n"
    "Local model CPU par slow hai, isliye summary background mein, queue, aur server start par auto warm-up. JavaScript se bane pages ke liye optional headless render add kiya, aur poori site crawl robots.txt, sitemap.xml, same-origin aur SSRF guard ke saath safe banaya.\n"
    "Future: bada ya fine-tuned model, docs par RAG, email alerts, CI integration, aur zyada checks."))
cols = [('Learning outcomes', EMERALD, ['Full-stack MERN with auth', 'JWT, validation, SSRF defence', 'Agent loop and tool calling', 'MCP: one registry, two clients', 'Testing with a mock LLM and a live one'], '✓'),
        ('Challenges solved', AMBER, ['A 3B model mangled an audit ID: short refs and title lookup', 'CPU LLM takes 15 to 40 s: background summary, queue, auto warm-up', 'Crawling a site safely: robots.txt, same-origin only, SSRF-guarded', 'Demo DB reset left a stale login ("Audit not found"): on-disk DB and session check'], '!'),
        ('Future enhancements', FUCHSIA, ['Larger or fine-tuned models', 'RAG over WCAG and OWASP docs', 'Email alerts for scheduled audits', 'CI integration (fail on low score)', 'More checks and Lighthouse-style metrics'], None)]
for i, (h, col, items, mark) in enumerate(cols):
    x = 0.7 + i * 4.05
    card(s, x, 1.95, 3.85, 4.85); shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.95, 3.85, 0.65, col, None, 0.25)
    text(s, x + 0.2, 1.95, 3.4, 0.65, (h, 18, True, WHITE), anchor=MSO_ANCHOR.MIDDLE)
    for k, t in enumerate(items):
        y = 2.8 + k * (0.78 if i != 1 else 0.98)
        if mark: text(s, x + 0.15, y, 0.4, 0.5, (mark, 16, True, col), anchor=MSO_ANCHOR.TOP)
        else: num(s, x + 0.18, y + 0.04, f'0{k + 1}', col, d=0.36, size=10.5)
        text(s, x + 0.6, y, 3.15, 0.75, (t, 14, False, INK))

# ================================================================ 12 Thank you
s = new('', '', '', (
    "Dhanyavaad. AuditPilot mein tools facts nikalte hain, local open-source LLM unhe explain karta hai, aur sab kuch aapke laptop par chalta hai.\n"
    "Ab main aapke sawalon ke liye taiyaar hoon.\n"
    "Agar aap chahein to main live demo dobara dikha sakta hoon, ya koi bhi feature detail mein explain kar sakta hoon."), dark=True)
text(s, 0, 1.6, 13.333, 0.4, ('END OF PRESENTATION', 13, True, 'F0ABFC'), PP_ALIGN.CENTER)
text(s, 0, 2.2, 13.333, 1.5, ('Thank you', 72, True, WHITE), PP_ALIGN.CENTER)
text(s, 0, 3.85, 13.333, 0.7, ('Learning. Building. Improving.', 26, False, 'F5D0FE'), PP_ALIGN.CENTER)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.2, 4.8, 0.93, 0.06, 'F0ABFC', None, 0.5)
text(s, 0, 5.1, 13.333, 1.0, [('Nikshey Yadav  ·  Roll No. 231302221', 18, True, WHITE), ('AuditPilot  ·  B.Tech CSE (Data Science)  ·  SGT University  ·  AccioJobs Pvt. Ltd.', 14, False, 'EDE9FE')], PP_ALIGN.CENTER)
text(s, 0, 6.6, 13.333, 0.4, ('REACT  ·  EXPRESS  ·  MONGODB  ·  OLLAMA', 12, True, 'F0ABFC'), PP_ALIGN.CENTER)

prs.save(OUT); print('saved', os.path.abspath(OUT), len(prs.slides._sldIdLst), 'slides')
